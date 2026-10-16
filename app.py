import streamlit as st
import time
import copy
import json
import os
import tempfile
import openpyxl
from openpyxl.styles import Alignment
import google.generativeai as genai

# ページの設定
st.set_page_config(page_title="小口現金 自動精算アプリ", layout="centered")

# --- UIカスタマイズ（ボタンを隠す） ---
st.markdown("""
    <style>
    [data-testid="stFileUploader"] section button {
        display: none;
    }
    </style>
""", unsafe_allow_html=True)

st.title("🧾 小口現金 自動精算アプリ")
st.write("レシート画像とマスターエクセルをドラッグ＆ドロップすると、AIが読み取ってエクセルに自動追記します。")

# --- APIキーの設定 ---
# StreamlitのSecretsからAPIキーを読み込んでGeminiを設定
try:
    genai.configure(api_key=st.secrets["GEMINI_API_KEY"])
except KeyError:
    st.error("APIキーが設定されていません。StreamlitのSecretsを確認してください。")

# --- 入力フォーム ---
st.subheader("1. 基本情報の入力")
col1, col2 = st.columns(2)
with col1:
    name = st.text_input("精算者氏名", placeholder="例：ベラジオ太郎")
with col2:
    allowance = st.number_input("今月支給額", min_value=0, step=1000, value=300000)

st.subheader("2. ファイルのアップロード")
uploaded_images = st.file_uploader(
    "レシート画像または動画をドラッグ＆ドロップ（複数可）", 
    accept_multiple_files=True
)

uploaded_excel = st.file_uploader("マスターエクセル（追記先）をドラッグ＆ドロップ")

# --- 実行ボタン ---
st.subheader("3. 処理の実行")
if st.button("データ抽出＆エクセル追記を開始"):
    if not name or not uploaded_images or not uploaded_excel:
        st.warning("氏名、レシート画像（動画）、マスターエクセルのすべてを入力・セットしてください。")
    else:
        with st.spinner("AIがレシートを解析し、エクセルに書き込んでいます...（数分かかる場合があります）"):
            try:
                # 1. 画像/動画をGeminiにアップロードして読み取らせる
                model = genai.GenerativeModel('gemini-1.5-flash')
                
                # プロンプト（AIへの複雑な指示ルール）
                prompt = """
                あなたは優秀な経理アシスタントです。添付されたレシート群を読み取り、以下のルールに従ってデータを抽出してください。
                【抽出ルール】
                1. 日付: ECサイトは「注文日」。形式は "YYYY/M/D"。
                2. 支払先: （）の補足削除、店舗名削除、法人格（株式会社等）削除。手書きは下部のゴム印・社名を正とする。
                3. 税率・インボイス: 1枚に8%と10%が混在する場合は2行に分ける。「T+13桁」があれば適格、なければ非適格とし、「適格10%」「適格8%」「非適格10%」「非適格8%」のいずれかに分類。
                4. 科目分類 (品目で判定): 
                   - 福利厚生費: 税率8%の飲食品
                   - 会議費: 税率10%の飲食代
                   - 接待交際費: 客単価5千円超の高級飲食、または贈答品
                   - 旅費交通費: ガソリン、駐車場、タクシー、公共交通
                   - 備品消耗品費: 物品購入（モノを買っている場合）
                   - 仮払金: 業務関連不明なもの
                5. 参加人数: 「会議費」または「接待交際費(飲食)」の場合のみ、金額(税込)を5500で割り、切り上げた整数（最低2名）を計算。それ以外は空文字にする。

                必ず以下のJSON配列のみを出力してください（Markdown記法は不要です）。
                [
                  {"date": "2026/10/16", "payee": "阪神百貨店", "tax_category": "適格10%", "amount": 1500, "item_details": "お茶", "category": "会議費", "participants": 2}
                ]
                """

                # 一時ファイルとして保存してAIに渡す準備
                uploaded_to_gemini = []
                for file in uploaded_images:
                    with tempfile.NamedTemporaryFile(delete=False, suffix="." + file.name.split(".")[-1]) as tmp:
                        tmp.write(file.getvalue())
                        tmp_path = tmp.name
                    
                    # Geminiにファイルをアップロード
                    gemini_file = genai.upload_file(tmp_path)
                    # 動画などの場合は処理完了を待つ
                    while gemini_file.state.name == "PROCESSING":
                        time.sleep(2)
                        gemini_file = genai.get_file(gemini_file.name)
                    uploaded_to_gemini.append(gemini_file)

                # AIに解析を実行させる
                response = model.generate_content([prompt] + uploaded_to_gemini)
                
                # 返ってきたJSONテキストを整形してPythonデータに変換
                json_text = response.text.strip().replace("```json", "").replace("```", "")
                extracted_data = json.loads(json_text)

                # 日付の古い順（昇順）に並べ替え
                extracted_data = sorted(extracted_data, key=lambda x: x["date"])

                # 2. エクセルファイルの処理（openpyxl）
                wb = openpyxl.load_workbook(uploaded_excel)
                source_sheet = wb.active
                # シートを複製
                ws = wb.copy_worksheet(source_sheet)
                ws.title = f"{name}_精算"

                # 【厳守事項】コピー元の書式や入力規則を復元
                ws.data_validations = copy.deepcopy(source_sheet.data_validations)
                ws.conditional_formatting._cf_rules = copy.deepcopy(source_sheet.conditional_formatting._cf_rules)

                # 基本情報の入力
                ws["F12"] = allowance

                # データ追記の開始行を探す（13行目以降でD列が空欄の最初の行を探す）
                start_row = 13
                while ws[f"D{start_row}"].value is not None:
                    start_row += 1

                # 抽出したデータを1行ずつエクセルに書き込む
                for i, item in enumerate(extracted_data):
                    r = start_row + i
                    ws[f"B{r}"] = name # 精算者氏名
                    ws[f"D{r}"] = item["date"]
                    ws[f"D{r}"].number_format = 'yyyy/m/d'
                    ws[f"D{r}"].alignment = Alignment(horizontal='right')
                    
                    # 以下、想定される列にデータを配置（※実際の列が違う場合はここを修正）
                    ws[f"C{r}"] = item.get("category", "")      # 科目
                    ws[f"E{r}"] = item.get("payee", "")         # 支払先
                    ws[f"F{r}"] = item.get("item_details", "")  # 品目・摘要
                    ws[f"G{r}"] = item.get("amount", "")        # 金額
                    ws[f"H{r}"] = item.get("participants", "")  # 参加人数
                    ws[f"J{r}"] = item.get("tax_category", "")  # 税率区分
                    
                    # 数式の維持: I列（会議費可能か？判定）
                    ws[f"I{r}"] = f'=IF(H{r}="","",IF(AND(H{r}>1,(G{r}/H{r})<=5500),"○","×"))'

                # 元のテンプレートシートを削除して完成
                wb.remove(source_sheet)

                # 完成したエクセルを一時ファイルに保存
                with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_out:
                    wb.save(tmp_out.name)
                    with open(tmp_out.name, "rb") as f:
                        processed_excel_data = f.read()

                st.success("✨ データの抽出とエクセルの追記がすべて完了しました！")
                
                # --- ダウンロードボタンの表示 ---
                st.download_button(
                    label="📥 完成したエクセルをダウンロード",
                    data=processed_excel_data,
                    file_name=f"処理済み_精算書_{name}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            except Exception as e:
                st.error(f"エラーが発生しました: {e}")
